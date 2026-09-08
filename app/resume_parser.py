from pathlib import Path
from docx import Document


def find_resume(resumes_folder: Path) -> Path | None:
    resumes = sorted(resumes_folder.glob("*.docx"))
    return resumes[0] if resumes else None


def extract_resume_text(resume_path: Path) -> str:
    document = Document(resume_path)
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

    return "\n".join(paragraphs)


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent.parent
    resume = find_resume(project_root / "resumes")

    if resume is None:
        print("No .docx resume found in the resumes folder.")
    else:
        print(extract_resume_text(resume)[:2000])
